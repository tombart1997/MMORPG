"""Build the offline illustrated docs and Excel workbook from sheets and CSVs."""
from pathlib import Path
import csv
import html
import json
import math
import re
import textwrap
import sys
import markdown
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.workbook.properties import CalcProperties
from build_balance import ART

BASE=Path(__file__).resolve().parents[1]
SHEETS=sorted((BASE/'sheets').glob('*.md'))
TOPICS=['Vision','Smithing','Equipment','Levels and Stats','Skill Mastery','Mob Bestiary','Loot Scaling','Fractures','World Regions','Classes and Combat','Quests and Town','Economy and Crafting','Mounts and Travel','UI and Feedback','Art Audio Production','Endgame and Online','Quality and Performance']
CSS='''
:root{--ink:#26362f;--muted:#687268;--jade:#244b3d;--gold:#a97c3e;--paper:#f6f2e8;--line:#d9d1c1}*{box-sizing:border-box}body{margin:0;background:var(--paper);color:var(--ink);font:16px/1.65 "Segoe UI",Arial,sans-serif}a{color:var(--jade);text-underline-offset:3px}a:hover{color:#966533}header{border-bottom:1px solid var(--line);padding:15px max(24px,calc((100vw - 1240px)/2));background:#fffaf2;display:flex;align-items:center;justify-content:space-between;gap:20px}header .brand{font:bold 22px Georgia,serif;letter-spacing:2px;text-decoration:none}header nav{display:flex;gap:22px;font-size:14px}main{max-width:1240px;padding:32px 26px 64px;margin:auto}.eyebrow{font-size:12px;text-transform:uppercase;letter-spacing:3px;color:var(--gold);font-weight:700}h1,h2,h3{font-family:Georgia,serif;line-height:1.2;color:#233d31}h1{font-size:44px;margin:12px 0 18px}h2{font-size:29px;margin:36px 0 16px}h3{font-size:22px;margin:24px 0 12px}.hero{display:grid;grid-template-columns:1.05fr 1fr;gap:36px;align-items:center}.hero img{width:100%;border-radius:5px;box-shadow:0 12px 35px #25382c20}.lead{font-size:18px;color:#556359}.actions{display:flex;gap:12px;flex-wrap:wrap;margin:24px 0}.button,button{background:var(--jade);border:1px solid var(--jade);color:white;padding:10px 18px;border-radius:4px;font:600 14px "Segoe UI",sans-serif;text-decoration:none;cursor:pointer}.button.secondary{background:transparent;color:var(--jade)}.facts{display:grid;grid-template-columns:repeat(4,1fr);border:1px solid var(--line);margin:30px 0;background:#fffcf5}.fact{padding:18px 24px;border-right:1px solid var(--line)}.fact:last-child{border:0}.fact strong{display:block;font:30px Georgia,serif}.fact span{font-size:13px;color:var(--muted)}.grid{display:grid;grid-template-columns:repeat(3,minmax(0,1fr));gap:20px}.tile{display:block;text-decoration:none;background:#fffcf5;border:1px solid var(--line);padding:22px;border-radius:4px;transition:transform .15s}.tile:hover{transform:translateY(-2px)}.tile h3{margin:8px 0 10px}.tile p{color:var(--muted);font-size:14px;margin:0}.tile img{width:100%;aspect-ratio:1.6;object-fit:cover;border-radius:3px}.note{padding:16px 20px;border-left:3px solid var(--gold);background:#ede6d6;font-size:14px}.doc{max-width:1030px}.doc img{max-width:100%;height:auto;display:block;border-radius:4px}.doc h1{font-size:38px}.doc p{max-width:92ch}.tablewrap{overflow-x:auto}table{border-collapse:collapse;width:100%;font-size:14px;margin:18px 0}th{background:#e2e9df;text-align:left;color:#284d3b}th,td{padding:10px 12px;border-bottom:1px solid var(--line);vertical-align:top}tbody tr:nth-child(even){background:#fdfaf3}code{background:#e7e6dc;border-radius:3px;padding:2px 5px;font-size:.88em}pre{overflow-x:auto;background:#263d31;color:#fcf8ea;padding:20px;border-radius:4px}pre code{background:none;color:inherit;padding:0}.controls{display:flex;flex-wrap:wrap;gap:18px;align-items:end;background:#fffcf5;border:1px solid var(--line);padding:20px;margin:24px 0}.controls label{font-size:13px;display:block;font-weight:600}.controls input,.controls select{font:16px "Segoe UI",sans-serif;padding:9px;border:1px solid #b4bbaa;border-radius:3px;background:white;color:var(--ink)}.controls input[type=number]{width:100px}.mob{border:1px solid var(--line);border-radius:5px;background:#fffcf6;overflow:hidden;align-self:start}.portrait{background-repeat:no-repeat;background-color:#efe0c1;aspect-ratio:1}.mob .content{padding:18px}.mob h2{font-size:23px;margin:0 0 8px}.mob .meta{color:var(--gold);font-size:13px;font-weight:600}.mob p{font-size:13px;color:var(--muted)}.mob table{font-size:12px}.mob th,.mob td{padding:7px 5px}.mob footer{padding:8px 18px;background:#eee8d9;font-size:12px}.count{font-size:14px;color:var(--muted)}[hidden]{display:none!important}.footer{margin-top:40px;padding-top:18px;border-top:1px solid var(--line);font-size:13px;color:var(--muted)}.gallery img{width:100%}@media(max-width:850px){.hero{grid-template-columns:1fr}.grid{grid-template-columns:repeat(2,minmax(0,1fr))}h1{font-size:35px}.facts{grid-template-columns:repeat(2,1fr)}header{padding:14px 20px}header nav{gap:12px}}@media(max-width:560px){.grid{grid-template-columns:1fr}.facts{grid-template-columns:repeat(2,1fr)}main{padding:22px 16px}header{display:block}header nav{margin-top:10px}h1{font-size:30px}}@media print{header,.actions,.controls{display:none}.grid{display:block}.tile,.mob{break-inside:avoid;margin:12px 0}body{background:white}main{max-width:none;padding:0}.doc img{max-height:170mm;object-fit:contain}.hero{display:block}.hero img{max-height:110mm;object-fit:contain}h2,h3{break-after:avoid}table{font-size:10px}th,td{padding:4px}.note{background:#eee}}
'''

def page(title,body,prefix=''):
    return f'''<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1"><title>{html.escape(title)} · Ironfield concepts</title><link rel="stylesheet" href="{prefix}concepts.css"></head><body><header><a class="brand" href="{prefix}index.html">IRONFIELD</a><nav aria-label="Main"><a href="{prefix}index.html">Design sheets</a><a href="{prefix}bestiary.html">Bestiary</a><a href="{prefix}art_gallery.html">Concept art</a><a href="{prefix}Ironfield_Concepts.xlsx">Workbook</a></nav></header>{body}</body></html>'''

def markdown_html(source):
    body=markdown.markdown(source,extensions=['tables','fenced_code','toc'])
    topic_files={path.name for path in SHEETS}
    def convert_link(match):
        target=match[1]
        fragment=match[2] or ''
        if Path(target).name in topic_files:
            target=target[:-3]+'.html'
        return f'href="{target}{fragment}"'
    body=re.sub(r'href="([^"#]+\.md)(#[^"]*)?"',convert_link,body)
    body=body.replace('<table>','<div class="tablewrap"><table>').replace('</table>','</table></div>')
    return body

def build_html():
    (BASE/'concepts.css').write_text(CSS,encoding='utf-8')
    tiles=[]
    snippets=[
        'The hunting, preparation and progression loop; fixed rules and first-release scope.',
        'Nine upgrade steps, exact odds and ingredients, destructive failure and accumulated risk.',
        'Equipment tiers, base-stat power, matte-to-shiny surfaces and blue/red/violet sparks.',
        'Quarter-level stat awards, XP thresholds and distinct VIT/STR/DEX/AGI effects.',
        '17 basic ranks, M1–M10 costs and G completion with 55 Mastery Leaves.',
        '37 creature designs with level, tier, EXP, full stats and individual drop rates.',
        'Symmetric logarithmic level-gap scaling, precise examples and reward scope.',
        'Destructible cores, defender waves, materials and the skill-mastery farming route.',
        'Briarwood, Cinderfall and future fen/heights, roads, landmarks and hunting bands.',
        'Four classes, eight paths, 48 skill designs and readable combat decisions.',
        'Town service roles, the current quest chain and introductory session targets.',
        'Exact trophy inputs, replacement gear, conversions, relic crafting and gold sinks.',
        'Horse anatomy, saddle/rider fit, unlocks, travel and render acceptance.',
        'XP milestones, allocation, mastery, destructive upgrade quotes and accessible feedback.',
        'Material language, export reviews, rendering budgets, audio and art provenance.',
        'Max-level goals, boss routes and proposed parties, trading, guilds and online authority.',
        'Diagnosis, bounded local repairs, opened comparisons, versioned replacement and actual game performance.'
    ]
    for i,path in enumerate(SHEETS):
        raw=path.read_text(encoding='utf-8')
        title=raw.splitlines()[0].removeprefix('# ')
        body=f'<main class="doc"><p class="eyebrow">Design sheet {i:02d} · 3 October 2026</p>{markdown_html(raw)}<p class="footer"><a href="../index.html">All design sheets</a> · <a href="{path.name}">Markdown source</a> · Fixed user rules with clearly marked proposed tuning.</p></main>'
        path.with_suffix('.html').write_text(page(title,body,'../'),encoding='utf-8')
        tiles.append(f'<a class="tile" href="sheets/{path.stem}.html"><span class="eyebrow">Sheet {i:02d}</span><h3>{TOPICS[i]}</h3><p>{snippets[i]}</p></a>')
    index=f'''<main><section class="hero"><div><p class="eyebrow">Game concept book · October 2026</p><h1>A dangerous road.<br>A stronger character.</h1><p class="lead">Ironfield's world, creatures and progression in separate illustrated design sheets.</p><div class="actions"><a class="button" href="sheets/01_smithing_upgrades.html">Smithing +9</a><a class="button secondary" href="bestiary.html">Explore the bestiary</a></div><p class="note">Your rules are fixed. New balance numbers and future features are marked as proposals. Current mob facts come from the game's resources.</p></div><img src="art/world_briarwood.png" alt="Briarwood frontier settlement and hunting roads"></section><section class="facts" aria-label="Contents"><div class="fact"><strong>{len(SHEETS)}</strong><span>Separate topic sheets</span></div><div class="fact"><strong>12</strong><span>Inspected concept plates</span></div><div class="fact"><strong>37</strong><span>Illustrated mob entries</span></div><div class="fact"><strong>151</strong><span>Individual item-drop rows</span></div></section><h2>Game systems and world</h2><div class="grid">{''.join(tiles)}</div><h2>Production assets</h2><div class="actions"><a class="button" href="../docs/full_asset_rebuild/gallery.html">All 111 replacement models</a><a class="button secondary" href="art/asset_turnarounds_v2/index.html">Practical multi-angle references</a></div><h2>Editable balance and art</h2><div class="actions"><a class="button" href="Ironfield_Concepts.xlsx">Open the Excel workbook</a><a class="button secondary" href="art_gallery.html">All concept art</a><a class="button secondary" href="README.md">Folder guide and rebuild instructions</a></div><p class="footer">Works offline. These are design documents and generated concept illustrations; gameplay changes are implementation targets.</p></main>'''
    (BASE/'index.html').write_text(page('Game concept book',index),encoding='utf-8')
    art=json.loads((BASE/'art/prompts.json').read_text(encoding='utf-8'))
    gallery=''.join(f'<section class="tile"><h2>{html.escape(j["topic"])}</h2><a href="art/{j["file"]}"><img src="art/{j["file"]}" alt="{html.escape(j["topic"])} concept plate" loading="lazy"></a><p>{j["file"]}</p></section>' for j in art)
    (BASE/'art_gallery.html').write_text(page('Concept art',f'<main class="gallery"><p class="eyebrow">Twelve production concept plates</p><h1>The art book</h1><p>Full-resolution generated illustrations, opened and inspected. <a href="art/ART_REVIEW.md">Art review</a> · <a href="art/prompts.json">Exact prompt set</a></p>{gallery}</main>'),encoding='utf-8')
    build_cards()
    print(f'Built offline index, {len(SHEETS)} illustrated HTML sheets, art gallery and interactive bestiary.')

def build_cards():
    data=json.loads((BASE/'tables/project_snapshot.json').read_text(encoding='utf-8'))
    cards=[]
    for m in sorted(data['enemies'],key=lambda m:(m['level'],m['tier'],m['id'])):
        plate,panel,brief=ART[m['id']]
        if plate=='bestiary_remaining_variants.png':
            # This painted board has a shorter upper row. CSS windows preserve
            # complete lower-row heads without saving or editing cropped PNGs.
            rectangles=[(.012,.012,.242,.422),(.255,.012,.485,.422),(.500,.012,.735,.422),(.755,.012,.990,.422),(.012,.405,.325,.970),(.350,.405,.638,.970),(.655,.405,.992,.970)]
            left,top,right,bottom=rectangles[panel-1]
        else:
            col=(panel-1)%3;row=(panel-1)//3
            # Small insets remove neighbouring silhouettes at painted grid edges.
            left=(col+.075)/3;right=(col+.925)/3
            top=(row+.06)/2;bottom=(row+.94)/2
        if plate=='bestiary_wolves_boars.png':
            rectangles=[(.015,.020,.310,.470),(.335,.000,.640,.470),(.670,.015,.990,.470),(.000,.505,.300,.985),(.308,.500,.645,.985),(.648,.500,.995,.985)]
            left,top,right,bottom=rectangles[panel-1]
        if m['id']=='young':left,top,right,bottom=.000,.005,.320,.458
        if m['id']=='ash_bandit':left,top,right,bottom=.012,.375,.325,.970
        width=right-left;height=bottom-top
        aspect=1536*width/(1024*height)
        style=f"background-image:url('art/{plate}');background-size:{100/width:g}% {100/height:g}%;background-position:{100*left/(1-width):g}% {100*top/(1-height):g}%;aspect-ratio:{aspect:g}"
        drops=[]
        for d in m['drops']:
            base=100*d['chance'];qty=str(d['min']) if d['min']==d['max'] else f"{d['min']}–{d['max']}"
            drops.append(f'<tr><td title="{html.escape(d["item_id"])}">{html.escape(d["name"])}</td><td>{base:g}%</td><td class="effective" data-base="{base}"></td><td>{qty}</td></tr>')
        cards.append(f'''<article class="mob" data-name="{html.escape(m['display_name'].lower())}" data-family="{m['family']}" data-level="{m['level']}"><a href="art/{plate}" aria-label="Open {html.escape(m['display_name'])} concept plate, panel {panel}"><div class="portrait" style="{style}"></div></a><div class="content"><h2>{html.escape(m['display_name'])}</h2><div class="meta">LEVEL {m['level']} · TIER {m['tier']} · {m['xp']} EXP</div><p>{html.escape(brief)}</p><p>{m['maximum_health']:g} HP · {m['defense']:g} Defense · {m['gold_min']}–{m['gold_max']} base gold</p><table><thead><tr><th>Drop</th><th>Base</th><th>At your level</th><th>Qty</th></tr></thead><tbody>{''.join(drops)}</tbody></table><p class="factor"></p></div><footer><a href="sheets/05_mob_bestiary.html#{m['display_name'].lower().replace(' ','-')}-{m['id']}">Full entry</a> · <a href="{html.escape('../'+m['source'])}">Resource</a> · panel {panel}</footer></article>''')
    script='''<script>const cards=[...document.querySelectorAll('.mob')];const level=document.getElementById('player-level');const search=document.getElementById('search');const family=document.getElementById('family');function refresh(){let visible=0;const p=Math.max(1,Math.min(100,Number(level.value)||1));for(const c of cards){const m=Number(c.dataset.level),gap=Math.abs(p-m),tol=Math.max(10,(p+m)/2);const f=Math.pow(Math.max(0,1-Math.log1p(gap)/Math.log1p(tol)),2);for(const cell of c.querySelectorAll('.effective')){const chance=Number(cell.dataset.base)*f;cell.textContent=chance===0?'0%':chance.toFixed(chance<.01?5:3)+'%';}c.querySelector('.factor').textContent='Proposed level multiplier: '+f.toFixed(4)+'× · gap '+gap;c.hidden=!(c.dataset.name.includes(search.value.toLowerCase())&&(family.value==='all'||c.dataset.family===family.value));if(!c.hidden)visible++;}document.getElementById('count').textContent=visible+' of '+cards.length+' creatures';}for(const el of [level,search,family])el.addEventListener('input',refresh);refresh();</script>'''
    families=sorted(set(m['family'] for m in data['enemies']))
    body=f'''<main><p class="eyebrow">Concept art + resource baseline</p><h1>The field bestiary</h1><p class="lead">37 creatures. Every entry includes its concept portrait, level, tier, EXP and full item-drop table.</p><p class="note">Base rates and EXP are current resource facts. “At your level” previews the proposed logarithmic item-drop penalty. Independent rolls can produce several items on one kill. Portraits are windows into the original concept plates; click to inspect the complete sheet.</p><div class="controls"><div><label for="search">Find a creature</label><input id="search" type="search" placeholder="Wolf, Spider, Warden…"></div><div><label for="family">Family</label><select id="family"><option value="all">All families</option>{''.join(f'<option value="{f}">{f.title()}</option>' for f in families)}</select></div><div><label for="player-level">Your level</label><input id="player-level" type="number" min="1" max="100" value="10"></div><span id="count" class="count" role="status"></span></div><div class="grid">{''.join(cards)}</div><p class="footer"><a href="sheets/06_loot_and_level_scaling.html">Formula and reward scope</a> · <a href="tables/mob_drops.csv">Editable item probabilities</a></p></main>{script}'''
    (BASE/'bestiary.html').write_text(page('Field bestiary',body),encoding='utf-8')

def parsed_value(value):
    if re.fullmatch(r'-?\d+',value):return int(value)
    if re.fullmatch(r'-?\d+\.\d+',value):return float(value)
    return value

def workbook():
    wb=Workbook();home=wb.active;home.title='Index';home.sheet_view.showGridLines=False
    home.append(['IRONFIELD — concepts and balance','Design baseline: 3 October 2026'])
    home.append(['Fixed user rules; proposed tuning is labelled. Current resource facts have their own status.'])
    home.append(['Open index.html for the illustrated browser edition.'])
    home.append(['Workbook changes are exploration only: they do not update CSV, Markdown or gameplay.'])
    home.append(['Topic','Workbook tab','Illustrated sheet'])
    home.column_dimensions['A'].width=68;home.column_dimensions['B'].width=35;home.column_dimensions['C'].width=65
    for i,path in enumerate(SHEETS):
        tab=f'{i:02d} {TOPICS[i]}'[:31]
        ws=wb.create_sheet(tab);ws.sheet_view.showGridLines=False;ws.sheet_view.zoomScale=85
        for col in range(1,9):ws.column_dimensions[get_column_letter(col)].width=12
        ws.column_dimensions['I'].width=4
        for col in range(10,19):ws.column_dimensions[get_column_letter(col)].width=11
        ws.freeze_panes='A3'
        raw=path.read_text(encoding='utf-8')
        row=1
        for line in raw.splitlines():
            if not line.strip() or line.startswith('![') or line.startswith('| ---') or line.startswith('```'):continue
            clean=re.sub(r'\[([^\]]+)\]\([^)]+\)',r'\1',line).replace('**','').replace('`','')
            heading=line.startswith('#')
            clean=clean.lstrip('# ').strip() if heading else clean
            if line.startswith('| '):
                clean='  |  '.join(part.strip() for part in clean.strip('|').split('|'))
            ws.merge_cells(start_row=row,start_column=1,end_row=row,end_column=8)
            cell=ws.cell(row,1,clean)
            cell.alignment=Alignment(wrap_text=True,vertical='top')
            cell.font=Font(name='Calibri',size=15 if heading else 11,bold=heading,color='244B3D' if heading else '283B31')
            if heading:cell.fill=PatternFill('solid',fgColor='E3EBDF')
            ws.row_dimensions[row].height=max(24 if heading else 20,15*math.ceil(len(clean)/100)+10)
            row+=1
        images=list(dict.fromkeys(re.findall(r'!\[[^\]]*\]\(\.\./art/([^)]*)\)',raw)))
        image_row=2
        for filename in images:
            picture=ExcelImage(BASE/'art'/filename)
            picture.height=picture.height*600/picture.width;picture.width=600
            ws.add_image(picture,f'J{image_row}')
            image_row+=int(picture.height/20)+5
        ws['J1']='Concept art — built-in image generation; full PNGs in art/'
        ws['J1'].font=Font(bold=True,color='A97C3E')
        ws.sheet_properties.pageSetUpPr.fitToPage=True
        ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A3
        ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
        home.append([TOPICS[i],tab,f'sheets/{path.stem}.html'])
        index_row=home.max_row
        home.cell(index_row,2).hyperlink=f"#'{tab}'!A1";home.cell(index_row,2).style='Hyperlink'
        home.cell(index_row,3).hyperlink=f'sheets/{path.stem}.html';home.cell(index_row,3).style='Hyperlink'
    table_names={'smithing_upgrades':'Smithing Table','gear_visuals':'Gear Visual Table','mobs':'Mob Stats Table','mob_drops':'Mob Drop Table','drop_scaling_examples':'Drop Calculator','level_progression':'XP Point Table','skill_mastery':'Mastery Cost Table','equipment_catalogue':'Equipment Table','fractures':'Fracture Table','fracture_materials':'Core Materials Table','class_skills':'48 Class Skills'}
    for filename,tab in table_names.items():
        path=BASE/'tables'/f'{filename}.csv'
        rows=list(csv.reader(path.open(encoding='utf-8-sig',newline='')))
        ws=wb.create_sheet(tab);ws.sheet_view.zoomScale=85
        for line in rows:ws.append([parsed_value(v) for v in line])
        headers=rows[0];cols={key:get_column_letter(i+1) for i,key in enumerate(headers)}
        ws.freeze_panes='D2';ws.row_dimensions[1].height=42
        for idx,header in enumerate(headers,1):
            width=max(14,min(45,max(len(str(r[idx-1])) for r in rows[:60])+2))
            ws.column_dimensions[get_column_letter(idx)].width=width
            ws.cell(1,idx).alignment=Alignment(wrap_text=True,vertical='center')
            ws.cell(1,idx).font=Font(bold=True,color='FFFFFF')
            ws.cell(1,idx).fill=PatternFill('solid',fgColor='244B3D')
        for row in ws.iter_rows(min_row=2):
            for c in row:c.alignment=Alignment(vertical='top',wrap_text=True)
            if filename=='class_skills':ws.row_dimensions[row[0].row].height=85
        table=Table(displayName=filename.replace('_','').title(),ref=f'A1:{get_column_letter(ws.max_column)}{ws.max_row}')
        table.tableStyleInfo=TableStyleInfo(name='TableStyleMedium4',showRowStripes=True)
        ws.add_table(table)
        if filename=='smithing_upgrades':
            for row in range(2,ws.max_row+1):
                c=cols
                ws[f'{c["failure_pct"]}{row}']=f'=100-{c["success_pct"]}{row}'
                first=2+((row-2)//9)*9
                ws[f'{c["route_survival_pct"]}{row}']=f'=100*PRODUCT({c["success_pct"]}{first}:{c["success_pct"]}{row})/100^({row-first+1})'
                ws[f'{c["expected_base_items_for_this_plus"]}{row}']=f'=100/{c["route_survival_pct"]}{row}'
            ws.conditional_formatting.add(f'{c["success_pct"]}2:{c["success_pct"]}{ws.max_row}',ColorScaleRule(start_type='num',start_value=15,start_color='DAA398',end_type='num',end_value=100,end_color='8DBC9B'))
        if filename=='drop_scaling_examples':
            c=cols
            for row in range(2,ws.max_row+1):
                p=f'{c["player_level"]}{row}';m=f'{c["mob_level"]}{row}';f=f'{c["multiplier"]}{row}'
                ws[f'{c["gap"]}{row}']=f'=ABS({p}-{m})'
                ws[f'{c["tolerance"]}{row}']=f'=MAX(10,({p}+{m})/2)'
                ws[f]=f'=MAX(0,1-LN(1+ABS({p}-{m}))/LN(1+MAX(10,({p}+{m})/2)))^2'
                ws[f'{c["trophy_effective_pct"]}{row}']=f'={c["trophy_base_pct"]}{row}*{f}'
                ws[f'{c["rare_gear_effective_pct"]}{row}']=f'={c["rare_gear_base_pct"]}{row}*{f}'
        ws.sheet_properties.pageSetUpPr.fitToPage=True
        ws.page_setup.orientation='landscape';ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
    for row in home.iter_rows():
        for cell in row:cell.alignment=Alignment(wrap_text=True,vertical='top')
    for row in range(1,home.max_row+1):home.row_dimensions[row].height=28
    home['A1'].font=Font(size=19,bold=True,color='244B3D')
    wb.calculation=CalcProperties(calcId=191029,fullCalcOnLoad=True)
    destination=BASE/'Ironfield_Concepts.xlsx'
    building=BASE/'.Ironfield_Concepts.building.xlsx'
    try:
        wb.save(building)
        building.replace(destination)
    finally:
        if building.exists():building.unlink()
    print(f'Built Excel workbook with {len(wb.sheetnames)} tabs and embedded original art.')

if __name__=='__main__':
    build_html()
    if '--html-only' not in sys.argv:
        workbook()
