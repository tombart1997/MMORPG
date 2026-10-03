"""Validate numerical design invariants, snapshot fidelity and all local links."""
from pathlib import Path
import csv
import json
import math
import re
from html.parser import HTMLParser
from urllib.parse import unquote,urlsplit
from PIL import Image
from openpyxl import load_workbook
from build_balance import CHANCES,ART,drop_factor,progress_at

BASE=Path(__file__).resolve().parents[1]
checks=0

def require(condition,message):
    global checks
    checks+=1
    if not condition:raise AssertionError(message)

def rows(name):
    return list(csv.DictReader((BASE/'tables'/name).open(encoding='utf-8-sig',newline='')))

class Links(HTMLParser):
    def __init__(self):super().__init__();self.links=[];self.ids=set()
    def handle_starttag(self,tag,attrs):
        attributes=dict(attrs)
        if 'id' in attributes:self.ids.add(attributes['id'])
        for key in ['href','src']:
            if key in attributes:self.links.append(attributes[key])

def validate():
    require(len(list((BASE/'sheets').glob('*.md')))==16,'16 topic sheets')
    require(len(list((BASE/'art').glob('*.png')))==12,'12 concept plates')
    up=rows('smithing_upgrades.csv');require(len(up)==45,'5 × 9 smithing steps')
    for tier in range(1,6):
        selected=[r for r in up if int(r['gear_tier'])==tier]
        require([int(r['to_plus']) for r in selected]==list(range(1,10)),f'Tier {tier}: complete steps')
        require(CHANCES[tier]==sorted(CHANCES[tier],reverse=True),f'Tier {tier}: higher + gets riskier')
        cumulative=1
        for r in selected:
            chance=float(r['success_pct']);step=int(r['to_plus']);cumulative*=chance/100
            require(chance+float(r['failure_pct'])==100,'Complementary probability')
            require(r['failure_result']=='DESTROY_ITEM_AND_SPEND_COSTS','Destructive failure')
            require(abs(float(r['route_survival_pct'])-100*cumulative)<1e-7,'Accumulated risk')
            require(r['spark_color']==('none' if step<7 else ['blue','red','violet'][step-7]),'Spark sequence')
            require(float(r['gold'])>0 and int(r['moss_shard'])==step,'Exact cost exists')
            if tier>1:require(chance<=CHANCES[tier-1][step-1],'Higher tier never easier')
    mastery=rows('skill_mastery.csv')
    require(len(mastery)==27,'17 basic + 10 mastery stages')
    require(mastery[16]['stage']=='M','Rank 17 becomes M')
    require(sum(int(r['mastery_leaf_cost']) for r in mastery)==55,'G costs 55 Leaves')
    require(mastery[-1]['result']=='G: Mastery II / maximum','M10 becomes G')
    require(float(mastery[-1]['potency_multiplier'])==1.62,'G potency')
    levels=rows('level_progression.csv');require(len(levels)==50,'50 proposed levels')
    for r in levels:
        l=int(r['level']);require(int(r['stat_points_earned_on_arrival'])==4*(l-1),'Four stat points per level')
        require(int(r['skill_points_earned_on_arrival'])==l,'One skill point per level including start')
    require([progress_at(1,x) for x in [0,149,150,299,300,449,450,599]]==[0,0,1,1,2,2,3,3],'Quarter thresholds exact')
    require(progress_at(2,0)==4,'Fourth point at level-up')
    require(progress_at(50,1000000)==196,'No repeatable max-level quarter rewards')
    require(progress_at(1,-100)==0,'Negative XP cannot remove earned points')
    for p in range(1,101):
        for m in range(1,101):
            f=drop_factor(p,m)
            require(0<=f<=1,'Bounded probability factor')
            require(f==drop_factor(m,p),'Symmetric scaling')
        require(drop_factor(p,p)==1,'Equal levels preserve base drops')
    require(drop_factor(25,5)==0 and drop_factor(5,25)==0,'Far-gap drops vanish in both directions')
    data=json.loads((BASE/'tables/project_snapshot.json').read_text(encoding='utf-8'))
    mobs=rows('mobs.csv');drops=rows('mob_drops.csv');require(len(mobs)==37,'37 unique mobs')
    require(set(r['mob_id'] for r in mobs)==set(ART),'Every mob has a unique concept panel')
    require(len(drops)==151,'151 current mob item rows')
    for m in data['enemies']:
        r=next(r for r in mobs if r['mob_id']==m['id'])
        require(int(r['level'])==m['level'] and int(r['tier'])==m['tier'] and int(r['exp'])==m['xp'],'Snapshot mob values')
        actual=[r for r in drops if r['mob_id']==m['id']]
        require(len(actual)==len(m['drops']),'Complete per-item drops')
        for d in m['drops']:
            r=next(r for r in actual if r['item_id']==d['item_id'])
            require(math.isclose(float(r['base_drop_pct']),100*d['chance'],abs_tol=1e-8),'Exact base drop chance')
    skills=rows('class_skills.csv');require(len(skills)==48,'48 class skills')
    require(len(set((r['base_class'],r['specialization']) for r in skills))==8,'Eight specialization designs')
    for image_path in (BASE/'art').glob('*.png'):
        with Image.open(image_path) as picture:
            require(picture.width>=1500 and picture.height>=850,'Full-size concept plate')
            picture.verify()
    parsers={}
    for path in BASE.rglob('*.html'):
        parser=Links();parser.feed(path.read_text(encoding='utf-8'));parsers[path.resolve()]=parser
    for path,parser in parsers.items():
        for link in parser.links:
            parsed=urlsplit(link)
            if parsed.scheme or link.startswith('//'):continue
            target=(path.parent/unquote(parsed.path)).resolve() if parsed.path else path
            require(target.exists(),f'Missing HTML link: {path.name} -> {link}')
            if parsed.fragment and target in parsers:
                require(unquote(parsed.fragment) in parsers[target].ids,f'Missing fragment: {link}')
    for path in BASE.rglob('*.md'):
        for link in re.findall(r'!?\[[^\]]*\]\(([^)]+)\)',path.read_text(encoding='utf-8')):
            parsed=urlsplit(link)
            if parsed.scheme or link.startswith('//'):continue
            require((path.parent/unquote(parsed.path)).resolve().exists(),f'Missing Markdown link: {path.name} -> {link}')
    for link in re.findall(r"background-image:url\('([^']+)'\)",(BASE/'bestiary.html').read_text(encoding='utf-8')):
        require((BASE/link).exists(),'Portrait image exists')
    wb=load_workbook(BASE/'Ironfield_Concepts.xlsx')
    require(len(wb.sheetnames)==28,'28 workbook tabs')
    require(len(wb['Mob Stats Table']['A'])==38,'Workbook contains all mobs')
    require(len(wb['Mob Drop Table']['A'])==152,'Workbook contains all drop rows')
    require(sum(len(ws._images) for ws in wb.worksheets)>=12,'Workbook embeds concept art')
    require(wb['Drop Calculator']['F2'].data_type=='f','Editable probability formula')
    require(wb['Smithing Table']['R2'].data_type=='f','Editable route risk formula')
    report={'checked_on':'2026-10-03','checks_passed':checks,'topic_sheets':16,'concept_plates':12,'unique_mobs':37,'mob_item_drops':151,'smithing_rows':45,'class_skill_designs':48,'workbook_tabs':28,'result':'PASS','scope':'Documentation consistency and numerical rules; no gameplay changes or runtime tests claimed.'}
    (BASE/'tables/validation_report.json').write_text(json.dumps(report,indent=2)+'\n',encoding='utf-8')
    print(f'PASS: {checks:,} design invariants and local links; images and workbook reopen successfully.')

if __name__=='__main__':
    validate()
